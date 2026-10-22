from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json, io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
REF = ROOT/'report/archive/template_reference/lab_01_template.docx'
OUT = ROOT/'report/archive/lab_02_report_draft_2026-10-20.docx'
DATA = ROOT/'results/three_controller_analysis'
ref_bytes = REF.read_bytes()
with ZipFile(io.BytesIO(ref_bytes)) as z:
    original = {n:z.read(n) for n in z.namelist()}
(ROOT/'report/archive/template_reference/inventory.json').write_text(json.dumps({
    'reference_sha256':hashlib.sha256(ref_bytes).hexdigest(),
    'parts':{n:{'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest()} for n,b in original.items()}},indent=2))
doc = Document(io.BytesIO(ref_bytes))
sect = deepcopy(doc.sections[0]._sectPr)
body = doc._element.body
for x in list(body): body.remove(x)
body.append(sect)
for name, size in [('Normal',11),('Title',19),('Heading 1',14),('Heading 2',12),('Heading 3',11),('Caption',9.5)]:
    if name not in doc.styles:
        doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    style=doc.styles[name]
    style.font.name='Aptos'
    style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(8)
    style.paragraph_format.line_spacing=278/240
    if name.startswith('Heading'):
        style.font.bold=True
        style.paragraph_format.keep_with_next=True
        style.paragraph_format.space_before=Pt(9)
doc.styles['Title'].font.bold=True
doc.styles['Normal'].paragraph_format.widow_control=True

def p(text): return doc.add_paragraph(text)
def h(text,level=1): return doc.add_paragraph(text,f'Heading {level}')
def page(): doc.add_page_break()
def eq(text):
    para=doc.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    math=OxmlElement('m:oMath'); run=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=text
    run.append(t); math.append(run); para._p.append(math)
    return para
def table(headers, rows, widths=None):
    tab=doc.add_table(rows=1,cols=len(headers)); tab.autofit=False
    available=(doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin)/914400
    widths=widths or [available/len(headers)]*len(headers)
    for col,w in zip(tab.columns,widths): col.width=Inches(w)
    for c,w in zip(tab.rows[0].cells,widths): c.width=Inches(w)
    for c,s in zip(tab.rows[0].cells,headers): c.text=str(s)
    rep=OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells=tab.add_row().cells
        for c,s,w in zip(cells,row,widths): c.text=str(s); c.width=Inches(w)
    for ri,row in enumerate(tab.rows):
        for cell in row.cells:
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');pr.append(va)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');pr.append(sh)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(1)
                par.paragraph_format.line_spacing=1.05
                for run in par.runs: run.font.size=Pt(9.5);run.bold=(ri==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return tab
def figure(filename,caption,width=5.75):
    par=doc.add_paragraph();par.paragraph_format.keep_with_next=True
    par.add_run().add_picture(str(DATA/filename),width=Inches(width))
    doc.add_paragraph(caption,'Caption')

doc.add_paragraph('Experiment II Look Ahead Car Following Control','Title')
p('INTR 5100 Traffic Flow Theory | Group 4 | Report draft | 20 October 2026')
p('We compare predictive PID, LQI and MPC controllers using three corrected physical runs and supplementary closed-loop simulations. MPC follows at smaller gaps, partly because it has a shorter target policy. PID attenuates middle-vehicle speed variation more in the two nonlaunch disturbance events. The evidence supports a trade-off between spacing policy and disturbance response, rather than a single overall winner.')
h('1 Experiment setup')
p('The experiment uses the SC²I scaled vehicle testbed and the supplied Shared Environment. The formation is leader → middle → student. The middle vehicle uses the supplied IDM implementation; the student receives current position and speed for both predecessors and can use their recorded histories. Its controller returns a desired speed to the main control program. Steering and the vehicle’s lower-level speed response are outside this callback.')
p('Only selected straight-road intervals are scored as student-controller operation, following the event-based approach used in Lab I. The program activates student control when middle and student occupy the same selected straight section; other intervals use the supplied bypass logic. Thus whole-lap duration and curve behavior cannot be assigned entirely to PID, LQI or MPC.')
table(['Setting','Recorded configuration'],[
('Vehicle geometry','Length 250 coordinate units; route loop 15,700.81 coordinate units'),
('Middle IDM defaults','v₀ 70; T 1.2; s₀ 250; a 20; b 30; δ 4; minimum spacing 80'),
('Straight selection','Three longest selected sections; observed event sequence 3, 1, 2, 3'),
('Control timing','Configured period 0.15 s; observed decision interval about 0.20 s'),
('Vehicle identity','Same physical student vehicle and same role identifiers across corrected runs'),
('Physical tests','One corrected 2.8 run each for PID, MPC and LQI; no physical 1× comparison')],[1.25,4.5])
p('Initial conditions differ, especially the leader–middle gap. The TA spacing shield is 80 units in the PID run and 50 in MPC/LQI; it never triggers in the scored intervals. The earlier run in which the TA inadvertently deployed Group 1’s controller is excluded. Controller labels follow the experiment records and group confirmation; deployed source hashes were not logged.')

page(); h('2 Controller design')
h('2.1 Geometry and scaling conventions',2)
p('Let q denote reported speed, v coordinate speed, u the returned speed command and s the forward bumper gap. Positions and gaps remain in map-coordinate units. We distinguish the telemetry conversion γ from the command response gain β, even when the submitted variants assign the same numerical value to both.')
eq('v = γq     s = forward route distance − 250     u = desired coordinate speed / β')
p('The 1× variant uses γ = β = 1 and the 2.8 variant uses γ = β = 2.8. This is a pair of lab hypotheses, not a generally established conversion. The cause of the discrepancy remains unresolved. Position-derived velocity is already in coordinate units per second and must not be multiplied by 2.8 again. Inside the middle-controller replica, coordinate speeds are divided by γ before invoking the supplied IDM formula; predicted motion is advanced in coordinate units. The returned command is not scaled again by the caller.')
h('2.2 Predictive PID and LQI',2)
p('The PID combines spacing feedback with a forecast-based speed reference. A bounded regression over recent speed history estimates leader acceleration. A 0.8 s rollout approximates the middle vehicle’s supplied IDM logic, rate limits, straight-road gating and response lag. Its inaccessible internal command memory is tracked approximately, so prediction is informative rather than exact.')
eq('s* = 300 + 1.8vₑ     e = s − s*     r = 0.55vₘ + 0.45v̂ₘ(t + 0.8)')
eq('u = clip and protect [(r + 0.12e + 0.002I + 0.08D) / β]')
p('I integrates spacing error and D is a filtered error derivative. Integral anti-windup, re-entry resets, stale-state checks, command rate limits and a closing-speed protection envelope limit aggressive behavior. LQI uses the same target policy and middle forecast, replacing the PID correction with a discrete linear quadratic integral feedback law. Its normalized states are spacing error, relative speed and accumulated error; Q = diag(4, 2, 1), with normalized control penalty 1. Neither the LQI gains nor the PID gains are retuned for this analysis.')
h('2.3 MPC and practical constraints',2)
p('The supplied MPC uses a 12-step horizon at 0.2 s, history-based predecessor trends and a 0.5 s assumed response lag. It penalizes gap error, predecessor-speed deviations and command changes while imposing speed, rate and predicted-gap constraints. Its target is s* = 250 + 1.2vₑ, shorter than PID/LQI. At q = 40 in the 2.8 version, the targets are 384.4 and 501.6 units respectively. Different own-target RMSEs therefore measure different policies.')
p('PID and LQI require standard-library imports plus the supplied route helper. MPC optimization additionally requires NumPy, CVXPY and OSQP; its internal fallback must be distinguished from optimization. Physical logs do not record this mode. A lab requirement for standard-library-only submissions therefore needs explicit attention before claiming the MPC optimizer is compliant. All submitted controller files remain unchanged.')

page(); h('2.4 Closed loop simulation results',2)
p('Each unchanged 2.8 controller is run independently against each recorded leader and middle input sequence. Simulated student position and speed are initialized once and subsequently evolve from that controller’s commands. Recorded student trajectories are used only for comparison, never injected into feedback. This is an ego closed-loop study with externally imposed front vehicles, rather than a fully endogenous three-vehicle simulation.')
eq('dvₑ/dt = [βplant u(t − d) − vₑ] / τ     dxₑ/dt = vₑ')
p('The plant gain, lag and delay are fitted separately for each recording. Actual decision timestamps, route gating and integer command rounding are reproduced; integration steps are at most 0.02 s. Telemetry uses the γ = 2.8 hypothesis. Nominal sensor inputs are used because measured obstacle detections cannot simply be transferred to a different simulated trajectory. The fitted plant and its open-loop validation are reported in Section 3.1.')
table(['Recorded input','PID RMSE','MPC RMSE','LQI RMSE'],[
('PID recording','137.8','79.5','146.0'),('MPC recording','113.0','69.4','119.9'),('LQI recording','120.5','71.0','133.8')])
p('Table 1. Baseline simulated own-target gap RMSE, in coordinate units. Each row shares recorded front inputs, initial student state and plant assumptions; target policies still differ.')
figure('closed_loop.png','Figure 1. Independent simulated student trajectories and gaps for the three input recordings. Black traces are physical references, not feedback states.',5.75)
p('Nine baseline tests plus half/double-lag sensitivity tests give 27 simulations. None records a sampled negative gap, TA spacing-shield activation or callback exception. MPC has the lowest own-target RMSE in the baseline model, followed by PID and LQI. This conditional ranking does not establish physical superiority: same-method simulated-versus-recorded student position RMSE is approximately 79–121 units, and the plant is only an exploratory approximation.')

page(); h('3 Experiment results analysis')
h('3.1 Open loop results and data preparation',2)
p('The physical tests ran with feedback. Here “open-loop results” means analysis using recorded signals without rerunning student feedback, including validation of a vehicle-response model driven by recorded commands. A controller replay on recorded student states could check command consistency, but could not predict the counterfactual vehicle response.')
p('We identify consecutive student_straight command intervals and verify that middle and student share a selected straight. All three trajectories are aligned on a 0.1 s grid without extrapolation. Samples whose source bracketing interval exceeds 0.3 s are rejected. Corresponding events are matched by section and occurrence, not by absolute elapsed time. Event 1 includes launch; event 4 is quiet and ends with the recording. No favorable initial-time cut is selected.')
p('Recorded rounded commands drive a first-order response model. Robust one-second route-displacement residuals are fitted on alternating 5 s blocks and evaluated on the other blocks. This is held-out time-block validation within a run, not independent-run validation. Model gains absorb unresolved scaling and dynamics and must not be interpreted as a revised 2.8 correction.')
table(['Input run','Gain','Lag s','Delay s','Held-out RMSE cu','Relative RMSE'],[
('PID','2.914','0.10','0.00','17.87','16.7%'),('MPC','2.937','0.10','0.10','17.64','16.5%'),('LQI','2.973','0.80','0.00','16.46','15.8%')],[.7,.65,.65,.7,1.55,1.5])
figure('plant_validation.png','Figure 2. Recorded and modeled one-second displacements. Validation error remains substantial; the fitted lag differs despite using the same student vehicle.',5.65)
p('PID and MPC select the lowest tested lag, indicating weak identification or omitted dynamics. Some source intervals are about 0.5 s long, and all runs show an unusually large position step near route coordinate 12,700 while reported speed remains 38. These features undermine unfiltered acceleration, jerk and position-derived TTC. They are flagged as measurement limitations rather than assigned to controller behavior.')

page(); h('3.2 Comparison between controllers',2)
p('Active decision counts are 337/768 for PID, 368/785 for MPC and 362/801 for LQI. Their active gap medians are 385.4, 293.9 and 425.7 units. These summaries exclude bypass intervals, but still mix launch and later events. The aligned event comparison below is more informative.')
table(['Event','Method','Entry → exit gap','Own RMSE','Bias','SD ego/mid'],[
('1','PID','116.4 → 411.4','114.5','−84.2','0.861'),('1','MPC','103.2 → 315.6','94.9','−84.1','0.898'),('1','LQI','113.9 → 458.1','107.5','−45.6','0.860'),
('2','PID','241.4 → 387.2','157.3','−147.6','0.866'),('2','MPC','146.9 → 262.2','175.2','−170.2','1.010'),('2','LQI','291.6 → 438.5','137.0','−119.8','1.002'),
('3','PID','518.6 → 477.1','57.2','46.1','0.776'),('3','MPC','372.6 → 379.4','37.4','30.2','0.834'),('3','LQI','664.6 → 463.6','131.6','102.3','0.795'),
('4','PID','414.0 → 455.5','55.6','−54.0','quiet'),('4','MPC','305.6 → 342.1','52.0','−50.3','quiet'),('4','LQI','401.1 → 468.3','61.3','−57.5','quiet')],[.45,.6,1.65,1,.85,1.2])
p('Table 2. Aligned straight-event metrics in coordinate units, except dimensionless reported-speed standard-deviation ratios. Positive bias means gap above the method’s own target. Ratios are withheld when middle-speed SD is below one raw unit.')
p('Event 2 exposes both recovery and initialization effects. LQI begins closer to its target than PID, while MPC begins at the smallest gap and has the largest own-target error. Event 3 is particularly informative: LQI removes about 201 units of inherited separation, whereas PID removes about 42 and MPC approximately holds its smaller target. A large LQI error is therefore not sufficient evidence of weak feedback.')
p('Matching the common route-coordinate coverage reduces coverage differences but does not equalize disturbances or entry state. On the common portion of event 3, median gaps are 491.4, 381.6 and 577.0 units; own-target RMSEs are 57.0, 38.3 and 131.9 for PID, MPC and LQI. The qualitative interpretation is unchanged.')
p('All three corrected runs show varying commands and bounded recorded gaps. MPC achieves closer following; PID shows stronger middle-to-ego attenuation in events 2 and 3; LQI recovers substantially from a large inherited gap. With one nonidentically initialized test per method, these are descriptive observations rather than statistically established controller rankings.')

page(); h('3.3 Microscopic efficiency comfort and safety',2)
p('Microscopic evaluation asks whether the student tracks its policy, recovers safely from disturbances and avoids excessive control activity. Own-target RMSE and bias quantify tracking. Gap reduction over an event quantifies recovery but depends strongly on the entry state. Whole-run times of 153.4, 156.9 and 160.1 s are recorded context only: curves, bypass control and initial formation also contribute.')
table(['Metric','PID','MPC','LQI'],[
('Active gap min / median / max','92.3 / 385.4 / 537.8','103.2 / 293.9 / 400.5','105.6 / 425.7 / 675.9'),
('Active spacing-shield flags','0','0','0'),
('Active obstacle-sensor flags','2','0','0'),
('Bypass obstacle-sensor flags','0','1','1'),
('RMS sent-command increment','1.48','1.50','1.46'),
('95th percentile absolute increment','4','3','3')],[2.05,1.23,1.23,1.24])
p('Table 3. Active-command overview. Command increments use consecutive active decisions and exclude pairs with an emergency flag. They are expressed in command units, not acceleration or jerk.')
h('Comfort and response timing',2)
p('RMS acceleration and jerk would normally characterize ride comfort. Numerical differentiation of the present coordinates produces spikes affected by localization and timing. Differentiating quantized reported speed has a different noise problem. Accordingly, command variation is reported as an actuator-demand proxy; the small differences in Table 3 do not establish a comfort winner. A future validated acceleration signal should support RMS and upper-percentile acceleration and jerk, with filter choice and sensitivity disclosed.')
p('Observed decision spacing is approximately 0.20 s. Previous local callback benchmarks place warmed PID/LQI calls around 1.1–1.4 ms at the 95th percentile and MPC around 1.8 ms, with slower initialization. These are host computation measurements, not physical end-to-end command latency. No communication or actuator timestamps are available to reconstruct that latency, and the closed-loop simulator does not advance time during callback computation.')
h('Safety interpretation',2)
p('No active recorded gap falls below the common 80-unit band, and no spacing-shield flags occur. This is evidence about these observations, not a collision-avoidance guarantee. PID’s two active sensor flags occur during launch; MPC and LQI each have one bypass sensor flag. A large route gap alone cannot establish that a sensor event is false.')
eq('TTC = s / (vₑ − vₘ)     when vₑ > vₘ')
p('TTC and required deceleration would complement minimum gap once speed and timing are validated. Under the present scale hypotheses they would be model-dependent, particularly near zero relative speed; no physical safety ranking is inferred from noisy derivative-based values. The observed minimum gaps also occur in different launch histories.')

page(); h('3.4 Macroscopic flow and string stability',2)
p('A three-vehicle experiment cannot directly establish a network fundamental diagram, queue discharge capacity or stability of an arbitrary-length platoon. It can reveal local headway behavior and how disturbances pass through this short formation. These two levels of evidence are kept distinct.')
eq('Rᵢⱼ = SD(qᵢ) / SD(qⱼ)')
p('R is a finite-window, demeaned speed-variation ratio on the same aligned samples. It is unchanged by a common constant 2.8 speed conversion. For an ideal linear single-predecessor chain, string stability is commonly studied by requiring transfer magnitude at most one across frequencies [1]. Our two-predecessor, gated and protected controller has multiple inputs and nonlinear switches; the observed R values neither estimate the full transfer matrix nor prove that condition.')
table(['Event','Method','Middle / leader','Ego / middle','Ego / leader'],[
(str(r['event']),r['method'],*(f"{r['gains'][k]:.3f}" for k in ['middle/leader','student/middle','student/leader']))
for event in range(1,4) for r in json.loads((DATA/'macro_metrics.json').read_text()) if r['event']==event],[.5,.7,1.5,1.5,1.55])
p('Table 4. Reported-speed SD ratios. Event 4 has insufficient excitation. Ratios are descriptive and have no confidence interval from repeated runs.')
p('In event 3, ego/middle is below one for all methods, but ego/leader is 1.061, 1.163 and 1.199 for PID, MPC and LQI. Thus the student damps variation relative to the middle vehicle without fully reversing the variation accumulated upstream. Event 2 shows PID attenuation at both stages, whereas MPC and LQI are near one at the student stage. Launch ratios are strongly affected by different starting speeds and transient duration.')
figure('macro_metrics.png','Figure 3. Local time gaps under γ = 2.8 and observed middle-to-student attenuation. Headways mix control response and inherited gaps.',5.75)
p('A stronger string-stability study would use repeated common disturbances, longer platoons, frequency-dependent gains, spacing-error amplification and peak response, with communication delay varied explicitly. The three current observations cannot be extrapolated to an arbitrary number of vehicles.')

page(); h('3.5 Local flow efficiency and fuel consumption',2)
h('Headway and flow proxies',3)
eq('hgap = s / vₑ     hspace = (s + 250) / vₑ     klocal = 1 / (s + 250)     Qlocal = vₑ / (s + 250)')
p('These quantities use vₑ = 2.8qₑ and exclude reported speeds below 5 raw units to avoid stopped-vehicle singularities. hgap is bumper time gap; hspace includes vehicle length. klocal is a reciprocal-spacing proxy, and Qlocal = klocal vₑ is its kinematic flow proxy, not measured detector flow or network throughput. Because metric coordinate conversion is unresolved, density is reported per 1,000 coordinate units rather than vehicles/km.')
macro=json.loads((DATA/'macro_metrics.json').read_text())
table(['Event','Method','Median hgap s','Density / 1000 cu','Flow proxy veh/s'],[
(str(r['event']),r['method'],f"{r['gap_time_median']:.2f}",f"{r['local_density_median']:.2f}",f"{r['local_flow_proxy_median']:.3f}")
for event in [2,3] for r in macro if r['event']==event],[.5,.7,1.4,1.6,1.55])
p('MPC’s smaller gaps yield higher local density and flow proxies in events 2 and 3. This is partly its shorter policy, rather than evidence that the feedback algorithm alone increases road capacity. LQI’s event-3 inherited gap lowers its density even while it closes the gap. A larger time gap can improve disturbance margins while reducing local space use; these objectives must be compared jointly.')
h('Fuel and energy demand',3)
p('Fuel consumption is an evaluation objective, but no fuel flow, battery current, voltage, mass or calibrated powertrain-efficiency map is logged. Reporting liters/100 km, Wh/km or a percentage fuel saving from these data would be unsupported. A suitable future measurement is integrated fuel flow per distance, or battery energy ∫VI dt per distance for an electric platform. A full-scale fuel estimate would require a specified reference vehicle and explicit mapping from lab motion to real motion.')
p('As a limited motion-based diagnostic, we compute positive kinetic variation per retained distance: K⁺ = Σmax[0, (v²ₖ − v²ₖ₋₁)/2] / ΣΔx. Adjacent retained 0.1 s pairs only are included; gaps in accepted data are never bridged. This measures repeated speed increases in coordinate units/s². It omits rolling resistance, aerodynamic drag, auxiliaries, efficiency and regeneration, and is sensitive to speed quantization; it is not fuel consumption or calibrated energy per mass.')
table(['Diagnostic','PID','MPC','LQI'],[
('K⁺ event 2',*(f"{next(r for r in macro if r['method']==m and r['event']==2)['positive_kinetic_variation_per_distance']:.2f}" for m in ['PID','MPC','LQI'])),
('K⁺ event 3',*(f"{next(r for r in macro if r['method']==m and r['event']==3)['positive_kinetic_variation_per_distance']:.2f}" for m in ['PID','MPC','LQI']))])
p('These values do not provide a consistent energy winner across events. A calibrated road-load model would additionally include acceleration, rolling resistance, aerodynamic drag, grade and powertrain losses, as in FASTSim [5]. Such a model should be validated against measured energy before claiming fuel benefits.')

page(); h('4 Discussion of look ahead design')
h('4.1 From a car following pair to two predecessors',2)
p('A conventional leader–follower pair uses the immediate predecessor’s motion and gap to select its own action. In our formation the middle–student pair is embedded behind an additional leader. The leader can reveal an approaching disturbance before it is fully visible in middle speed. Because the middle uses known IDM logic, leader–middle gap and relative speed help predict how it will respond. Cooperative control literature motivates combining local feedback with upstream feedforward rather than waiting for the entire disturbance to reach the nearest vehicle [1–3].')
p('Knowing the IDM equation does not make middle motion exactly predictable. Hidden command memory, integer rounding, main-program gating, sensor stops, actuation lag and asynchronous packets separate a mathematical IDM acceleration from measured vehicle motion. History helps estimate trends and internal state, while a short horizon and blended reference reduce dependence on an uncertain forecast. A predictor should reproduce the actual supplied implementation before adding sophistication.')
h('4.2 Why attenuation and spacing accuracy can conflict',2)
p('When the leader begins to slow, predicted middle speed may fall before the middle actually slows. The student then reduces its reference earlier. This can suppress a later speed transient but temporarily enlarge its gap and increase spacing error. The feedback and integral terms must recover that gap; protections and command-rate limits can delay recovery. A smoother speed trace and a larger transient gap can therefore arise from the same anticipatory action. They measure different objectives.')
p('The explanation is a mechanism, not a causal finding from the three physical tests: every tested controller uses look-ahead information, and there is no matched physical pair-only baseline. LQI’s large event-3 gap was already present on entry. MPC also has a shorter target policy. These facts prevent attributing every observed gap difference to prediction or controller class.')
h('4.3 A fair pair model comparison',2)
p('A useful follow-up is an ablation with the same PID gains, target, safeguards, plant, initial state and disturbances. Compare immediate-middle-speed reference only; middle-history trend prediction; and leader-assisted IDM prediction. Retain both γ = β = 1 and γ = β = 2.8 versions, and vary telemetry and command gains independently in sensitivity tests. Equalize the target policy when comparing PID, LQI and MPC. This would separate the benefit of information from the benefit of a different desired headway.')
p('Score own and common-policy spacing errors, peak gap deviation, response delay, minimum gap, validated comfort measures, full-chain and adjacent disturbance amplification, and calibrated energy per distance. Include the cost of stale or delayed upstream information. For the macroscopic claim, repeat with more vehicles and common disturbances; for fuel claims, measure energy or specify and validate a powertrain model. No controller revisions are made as part of this report analysis.')
p('The current evidence supports PID as a promising disturbance-attenuation option and MPC as a closer-following option under its chosen policy. LQI demonstrates substantial recovery from a large entry gap. More controlled repetitions are required to select a method across efficiency, safety, comfort, stability and energy objectives.')

page(); h('References')
refs=[
'[1] G. J. L. Naus, R. P. A. Vugts, J. Ploeg, M. J. G. van de Molengraft and M. Steinbuch. Cooperative adaptive cruise control, design and experiments. American Control Conference, 2010, pp. 6145–6150. https://doi.org/10.1109/ACC.2010.5531596',
'[2] C. Wang, S. Gong, A. Zhou, T. Li and S. Peeta. Cooperative adaptive cruise control for connected autonomous vehicles by factoring communication-related constraints. Transportation Research Part C, 113, 2020, pp. 124–145. https://doi.org/10.1016/j.trc.2019.04.010',
'[3] C. Wu, Y. Lin and A. Eskandarian. Cooperative Adaptive Cruise Control With Adaptive Kalman Filter Subject to Temporary Communication Loss. IEEE Access, 2019. https://doi.org/10.1109/ACCESS.2019.2928004',
'[4] M. R. Cantas, S. Y. Gelbal, L. Guvenc and B. A. Guvenc. Cooperative Adaptive Cruise Control Design and Implementation. SAE Technical Paper 2019-01-0496. Author manuscript: https://arxiv.org/abs/2306.08182',
'[5] A. Brooker, J. Gonder, L. Wang, E. Wood, S. Lopp and L. Ramroth. FASTSim A Model to Estimate Vehicle Efficiency, Cost and Performance. SAE Technical Paper 2015-01-0973, 2015. National laboratory publication record: https://research-hub.nlr.gov/en/publications/fastsim-a-model-to-estimate-vehicle-efficiency-cost-and-performan-2/',
]
for text in refs: p(text)
h('Experimental sources and reproducibility',2)
p('The Shared Environment and Lab II guideline define the interface and operating conditions. The Lab I report supplies the straight-event processing precedent, while the Lab I report template supplies page geometry and typography. The tutor’s four-part framework governs this report’s structure.')
p('Corrected archives are preserved under data/raw/group4_pid_2p8, group4_mpc_2p8 and group4_lqi_2p8, each with an archive manifest. Primary analysis outputs are physical_summary.json, closed_loop.json, aligned event CSV files and macro_metrics.json in results/three_controller_analysis. These include all four straight events and the 27 supplementary simulations. Figures are generated from these local outputs.')
p('Reproduction entry points are src/report_three_runs.py, src/recorded_closed_loop.py and src/report_macro.py in the Lab II analysis environment. Analysis dependencies are separate from controller submission requirements. Archived controller and Shared Environment files are not changed by these analyses.')

temp=io.BytesIO();doc.save(temp)
with ZipFile(io.BytesIO(temp.getvalue())) as z: built={n:z.read(n) for n in z.namelist()}
# Preserve opaque template parts byte-for-byte; only body, styles and new media/relationships change.
editable={'word/document.xml','word/styles.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
for n,b in original.items():
    if n not in editable: built[n]=b
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
    for n,b in built.items(): z.writestr(n,b)
assert REF.read_bytes()==ref_bytes
with ZipFile(OUT) as z:
    for n,b in original.items():
        if n not in editable: assert z.read(n)==b,n
print(OUT)
