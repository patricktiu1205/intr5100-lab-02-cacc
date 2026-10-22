"""Report-only revision from immutable experimental and saved analysis records.

Use the shared intr5100 Conda environment. No controller, fitting or simulation
module is imported. Both editable formats use the same content blocks.
"""
from pathlib import Path
from copy import deepcopy
from hashlib import sha256
from zipfile import ZipFile, ZIP_DEFLATED
import csv
import io
import json
import os
import argparse
import shutil
import subprocess
import sys
import tempfile

import numpy as np
from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'results/three_controller_analysis'
FIG = ROOT / 'report/source/figures'
METHODS = ('PID', 'MPC', 'LQI')


def protected():
    folders = ['data', 'controllers', 'environment', 'results', 'simulation_results', 'src']
    files = [p for folder in folders for p in (ROOT / folder).rglob('*') if p.is_file()]
    return {str(p.relative_to(ROOT)): sha256(p.read_bytes()).hexdigest()
            for p in sorted(files) if '__pycache__' not in p.parts and p.name != '.DS_Store'}


def rows(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--skip-pdf', action='store_true', help='Build DOCX/LaTeX without invoking Tectonic')
    parser.add_argument('--offline', action='store_true', help='Use cached Tectonic resources only')
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='intr5100-lab2-report-') as cache:
        os.environ.setdefault('MPLCONFIGDIR', cache)
        before = protected()
        physical = json.loads((DATA/'physical_summary.json').read_text())
        closed = json.loads((DATA/'closed_loop.json').read_text())
        macro = json.loads((DATA/'macro_metrics.json').read_text())
        checks = []
        for method in METHODS:
            raw = ROOT/f'data/raw/group4_{method.lower()}_2p8'
            manifest = json.loads((raw/'archive_manifest.json').read_text())
            for item in manifest['files']:
                assert sha256((raw/item['file']).read_bytes()).hexdigest() == item['sha256']
            commands = [r for r in rows(raw/'commands.csv') if r['role'] == 'student']
            active = [r for r in commands if r['mode'] == 'student_straight']
            summary = physical['physical'][method]
            assert len(active) == summary['active_calls']
            assert len(commands) == summary['total_calls']
            for label, aggregate in [('min', min), ('median', np.median), ('max', max)]:
                assert np.isclose(aggregate([float(r['spacing']) for r in active]), summary[f'active_gap_{label}'])
            for event in summary['events']:
                rr = rows(DATA/f"{method.lower()}_event_{event['event']}.csv")
                values = {k: np.array([float(r[k]) for r in rr]) for k in rr[0]}
                assert len(rr) == event['samples']
                assert np.isclose(np.sqrt(np.mean(values['error']**2)), event['own_rmse'])
                if event['reported_speed_ratio'] is not None:
                    assert np.isclose(np.std(values['student_v'])/np.std(values['middle_v']), event['reported_speed_ratio'])
            validation = [r for r in rows(DATA/f'plant_validation_{method.lower()}.csv')
                          if r['valid'] == '1' and r['training'] == '0']
            rmse = np.sqrt(np.mean([(float(r['predicted_displacement'])-float(r['observed_displacement']))**2
                                    for r in validation]))
            assert np.isclose(rmse, closed['plant_fits'][method]['validation_displacement_rmse'])
            checks.append(f'{method}: archive hashes, active gap summaries, four event errors/ratios, saved plant validation RMSE')

        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        plt.rcParams.update({'font.size': 9, 'axes.spines.top': False, 'axes.spines.right': False})
        FIG.mkdir(parents=True, exist_ok=True)
        fig, axes = plt.subplots(3, 2, figsize=(7.0, 5.5), sharex='col')
        for i, method in enumerate(METHODS):
            r = rows(DATA/f'{method.lower()}_event_3.csv')
            t = np.array([float(x['t']) for x in r]); t -= t[0]
            breaks = np.flatnonzero(np.diff(t) > .15)+1
            x = np.insert(t, breaks, np.nan)
            channel = lambda k: np.insert(np.array([float(z[k]) for z in r]), breaks, np.nan)
            axes[i, 0].plot(x, channel('gap'), color='#2166ac', label='Recorded gap')
            axes[i, 0].plot(x, channel('target'), color='#b35806', ls='--', label='Own target')
            for key, color, style, label in [('leader_v', '#777777', ':', 'Leader'),
                                            ('middle_v', '#b35806', '--', 'Middle'),
                                            ('student_v', '#2166ac', '-', 'Student')]:
                axes[i, 1].plot(x, channel(key), color=color, ls=style, label=label)
            axes[i, 0].set_ylabel(f'{method}\nGap (cu)')
            axes[i, 1].set_ylabel('Reported speed')
            for ax in axes[i]:
                ax.grid(alpha=.2)
        axes[0, 0].legend(fontsize=7, ncol=2)
        axes[0, 1].legend(fontsize=7, ncol=3)
        for ax in axes[-1]: ax.set_xlabel('Seconds since each event entry')
        fig.tight_layout(pad=.7)
        for ext in ('png', 'pdf'): fig.savefig(FIG/f'event_3_records.{ext}', dpi=200)
        plt.close(fig)

        # One content model for both formats. Numbered sections follow the original
        # setup/design/results/discussion outline; page breaks keep evidence together.
        blocks = []
        def p(s): blocks.append(('p', s))
        def h(s, level=1): blocks.append(('h', level, s))
        def page(): blocks.append(('page',))
        def table(caption, headers, values, widths): blocks.append(('table', caption, headers, values, widths))
        def equation(tex, text): blocks.append(('equation', tex, text))

        blocks.append(('title', 'Look Ahead Car Following Control'))
        blocks.append(('subtitle', 'INTR 5100   Lab Session II   Group 4'))
        h('Overview', 2)
        p('We compare predictive PID, LQI and MPC using one recorded Group 4 run per method and previously saved simulations. The MPC-labelled run follows at smaller gaps under a shorter target policy. PID has lower student-to-middle speed variation ratios in the two excited nonlaunch events. LQI closes a large inherited gap. These observations show different spacing and disturbance responses without establishing an overall controller ranking.')
        h('1 Experiment setup and data preparation')
        p('The SC2I testbed uses a prescribed-speed leader, an IDM-controlled middle vehicle and a student-controlled follower. The student callback receives all three states and their histories, then returns a desired speed command. Steering and the lower-level speed response remain in the supplied program. Student control operates only when middle and student share a selected straight; bypass intervals use the supplied logic.')
        table('Recorded experimental configuration', ['Setting', 'Value'], [
            ['Geometry', 'Vehicle length 250 cu; route loop 15,700.81 cu'],
            ['Middle IDM', 'v0 70; T 1.2; s0 250; a 20; b 30; exponent 4'],
            ['Timing', 'Configured 0.15 s; observed decisions about 0.20 s'],
            ['Student spacing shield', 'PID 80 cu; MPC and LQI 50 cu'],
            ['Physical data', 'One 2.8-labelled run per method; same role identifiers']
        ], [1.35, 5.35])
        p('The earlier Group 1 deployment-error run is excluded. Initial gaps and predecessor histories differ across the retained runs. Source hashes and MPC solver status were not logged during physical testing; method names identify the recorded runs, rather than proving which internal control mode produced every command.')
        h('Event processing', 2)
        p('Consecutive student_straight decisions define four events, matched by selected-section index and occurrence: 3, 1, 2, 3. The three trajectories are aligned on a 0.1 s grid without extrapolation; source interpolation spans above 0.3 s are rejected. Middle and student must share the same selected section. Event 1 includes launch; event 4 is quiet and ends with the recording.')
        p('Spacing is forward route distance minus 250 cu. Position jumps and some approximately 0.5 s telemetry gaps limit differentiated acceleration, jerk and TTC. Coordinate units are retained because physical scale and command-response gain remain unresolved.')

        page(); h('2 Controller design')
        h('Speed conventions and target policies', 2)
        p('Let q be reported speed, v coordinate speed and u the returned command. Telemetry conversion gamma and command response beta are distinct hypotheses. The 2.8 variants set both to 2.8; this does not establish either as a measured platform constant. Position-derived speeds already use cu/s and receive no additional factor.')
        equation(r'v=\gamma q,\qquad u=v_{\rm desired}/\beta,\qquad \gamma=\beta=2.8',
                 'v = γq     u = desired coordinate speed / β     γ = β = 2.8')
        table('Own target bumper gaps in the 2.8 variants', ['Method', 'Target gap in cu', 'At q = 40'], [
            ['PID and LQI', '300 + 1.8v', '501.6 cu'], ['MPC', '250 + 1.2v', '384.4 cu']
        ], [1.35, 3.7, 1.65])
        p('MPC therefore seeks a gap 117.2 cu shorter at reported speed 40. Own-target error evaluates each policy and cannot by itself rank all methods against a common spacing objective.')
        h('Predictive PID and LQI', 2)
        p('PID estimates leader acceleration from recent history and uses a 0.8 s rollout of the supplied middle IDM logic with an assumed 0.5 s response lag. It blends current and forecast middle speed, then adds spacing feedback. The middle command memory is approximated; forecast accuracy is limited by gating, sensor stops and vehicle response.')
        equation(r'e=s-(300+1.8v),\qquad r=0.55v_M+0.45\widehat v_M(t+0.8)',
                 'e = s − (300 + 1.8v)     r = 0.55vM + 0.45 forecast vM')
        equation(r'u_{\rm nominal}=(r+0.12e+0.002I+0.08D)/\beta',
                 'Nominal PID command = (r + 0.12e + 0.002I + 0.08D) / β')
        p('I is accumulated spacing error and D its filtered derivative. Anti-windup, inactive-interval resets, stale-state checks, command limits and closing protection modify the nominal command. LQI uses the same target and forecast with normalized spacing, relative-speed and integral states. Its discrete Riccati feedback uses Q = diag(4, 2, 1) and unit control penalty. These designs and gains are retained unchanged.')
        h('MPC and implementation limits', 2)
        p('MPC uses 12 steps of 0.2 s, history-based predecessor trends and an assumed 0.5 s lag. Its objective penalizes gap error, speed deviations and command changes, subject to speed, rate and predicted-gap constraints. Optimization requires NumPy, CVXPY and OSQP; otherwise its feedback fallback may operate. Physical logs do not distinguish that fallback from successful optimization.')

        page(); h('3 Recorded results and analysis')
        h('3.1 Active control and spacing errors', 2)
        values=[]
        for m in METHODS:
            s=physical['physical'][m]
            values.append([m, f"{s['active_calls']} / {s['total_calls']}",
                           f"{s['active_gap_min']:.1f}", f"{s['active_gap_median']:.1f}",
                           str(s['active_spacing_flags']), str(s['active_sensor_flags'])])
        table('Student command-log overview during active control',
              ['Run', 'Active / total', 'Min gap', 'Median gap', 'Spacing flags', 'Sensor flags'], values,
              [.65, 1.3, 1.05, 1.15, 1.25, 1.3])
        p('Gaps are in cu and flag counts are command rows, not independent incidents. All active gaps exceed a common 80-cu band. PID has two consecutive active sensor-flag rows; MPC and LQI each have one sensor flag during bypass. No active spacing shield or wrapper fallback is recorded. Sensor-event causes and MPC internal fallback status remain unverified.')
        values=[]
        for i in range(4):
            for m in METHODS:
                e=physical['physical'][m]['events'][i]
                values.append([str(i+1),m,str(e['samples']),f"{e['entry_gap']:.1f} to {e['exit_gap']:.1f}",
                               f"{e['own_rmse']:.1f}",f"{e['bias']:.1f}"])
        table('Aligned trajectory samples and own-target spacing errors',
              ['Event', 'Run', 'Samples', 'Entry to exit gap', 'RMSE', 'Bias'], values,
              [.6,.6,.8,2.25,1.2,1.25])
        p('RMSE is the square root of mean squared gap error; bias is mean gap minus target. Both use retained aligned samples and cu. Positive bias denotes a gap above the method\'s target. These event metrics differ from command-log summaries because their sampling and synchronization masks differ.')
        p('Event 2 begins below target for every method, with different entry gaps. In event 3, LQI removes about 201 cu of separation, PID about 42 cu, while MPC approximately maintains its smaller gap. LQI\'s large inherited gap contributes to its higher RMSE. A common-route restriction preserves this interpretation without equalizing initial states or disturbance histories.')

        page(); h('3.2 Trajectories and disturbance propagation', 2)
        blocks.append(('figure', 'event_3_records', 'Event 3 recorded gaps, own targets and all three reported speeds. Each row starts at its own event entry; elapsed time does not represent matched disturbances across runs.'))
        values=[]
        for i in (2,3):
            for m in METHODS:
                rec=next(r for r in macro if r['method']==m and r['event']==i)
                g=rec['gains']
                values.append([str(i),m,f"{g['middle/leader']:.3f}",f"{g['student/middle']:.3f}",f"{g['student/leader']:.3f}"])
        table('Demeaned reported-speed standard deviation ratios',
              ['Event','Run','Middle / leader','Student / middle','Student / leader'], values,
              [.65,.65,1.8,1.8,1.8])
        p('Ratios use the same retained samples within each event; a common constant speed conversion cancels. PID has the smallest student/middle ratio in events 2 and 3. In event 3 all methods attenuate middle-speed variation, yet student/leader remains above one, so upstream amplification is not fully removed. Event 4 has zero middle-speed variation and no meaningful ratio.')
        p('These finite-window observations do not prove string stability. The formation has only three vehicles, unequal inputs and switched protections; frequency-dependent amplification and longer-platoon behavior are not established [1].')

        page(); h('3.3 Previously saved simulation evidence', 2)
        p('The following numbers come from the existing analysis records. No response model was refitted and no simulation was repeated for this revision. These supplementary results are kept separate from the physical measurements.')
        h('Recorded command response model', 2)
        p('The saved first-order plant fit used rounded commands and one-second route-displacement residuals, with alternating 5 s training and validation blocks. This is validation within one recording, rather than on an independent experimental run. The gain absorbs unresolved scale and dynamics.')
        values=[]
        for m in METHODS:
            f=closed['plant_fits'][m]
            values.append([m,f"{f['beta']:.3f}",f"{f['lag']:.2f}",f"{f['delay']:.2f}",
                           f"{f['validation_displacement_rmse']:.2f}",f"{100*f['validation_relative_rmse']:.1f}%"])
        table('Saved held-out one-second displacement errors',
              ['Input run','Gain','Lag s','Delay s','RMSE cu','Relative RMSE'], values,
              [1.15,.85,.85,.9,1.25,1.7])
        p('Relative RMSE divides by mean held-out displacement. PID and MPC select the lowest tested lag, 0.10 s, indicating weak identification or omitted dynamics. The differing fits for the same student vehicle and 15.8-16.7% errors limit the model\'s predictive authority.')
        h('Student feedback with recorded predecessors', 2)
        p('In the saved study, each controller drove a simulated student using recorded leader and middle trajectories. Student position and speed evolved from commands after initialization; recorded student states were reference data. Front vehicles remained externally imposed. Assumptions included the 2.8 telemetry scale, source-specific plant fit, actual decision times, bypass gating and nominal obstacle inputs.')
        values=[]
        for source in METHODS:
            values.append([source]+[f"{next(r for r in closed['runs'] if r['source']==source and r['controller']==m and r['lag_multiplier']==1)['own_rmse']:.1f}" for m in METHODS])
        table('Saved baseline own-target gap RMSE in cu', ['Recorded front input','PID','MPC','LQI'], values,
              [2.8,1.3,1.3,1.3])
        p('Nine baseline cases plus half/double-lag variants total 27 saved simulations. They record no sampled negative gap, spacing-shield activation or callback exception. Some MPC decisions are protective stops, so even the simulated MPC cases are not optimization-only trajectories. These sampled outcomes provide conditional evidence, not a collision-avoidance guarantee.')
        p('MPC has the lowest baseline own-target RMSE under this model. Different target policies and approximately 79-121 cu same-method simulated-versus-recorded position RMSE prevent translating that ranking into demonstrated physical superiority.')

        page(); h('3.4 Efficiency comfort and safety', 2)
        h('Local spacing and flow', 2)
        p('For moving samples with reported speed at least 5, use v = 2.8q, bumper time gap s/v, reciprocal space occupancy 1000/(s + 250), and local flow proxy v/(s + 250). The latter quantities describe one pair\'s local spacing and motion; they are not measured detector flow, network density or road capacity.')
        values=[]
        for i in (2,3):
            for m in METHODS:
                r=next(r for r in macro if r['event']==i and r['method']==m)
                values.append([str(i),m,f"{r['gap_time_median']:.2f}",f"{r['local_density_median']:.2f}",f"{r['local_flow_proxy_median']:.3f}"])
        table('Median kinematic proxies under the 2.8 speed hypothesis',
              ['Event','Run','Time gap s','Vehicles / 1000 cu','Flow proxy veh/s'],values,
              [.65,.65,1.5,2.05,1.85])
        p('MPC has higher local flow proxies in these events, consistent with its shorter gap policy. Inherited separation also affects the values. These data cannot isolate an algorithmic improvement in network throughput, and unresolved metric scale prevents conversion to vehicles/km.')
        h('Comfort and timing', 2)
        p('The RMS sent-command increments are PID 1.48, MPC 1.50 and LQI 1.46 command units. They use consecutive active decisions excluding flagged pairs. These small differences describe command activity, not measured acceleration, jerk or ride comfort. Localization steps and speed quantization make a physical comfort ranking unreliable.')
        p('The approximately 0.20 s decision interval is loop cadence, not end-to-end response latency. Previously saved local callback benchmarks exclude communication and actuator delay and do not establish physical timing performance.')
        h('Safety and energy', 2)
        p('Observed active minimum gaps are 92.3, 103.2 and 105.6 cu for PID, MPC and LQI. None activates the recorded spacing shield, although native thresholds differ. Unequal launch gaps, unverified sensor causes and limited repetition prevent a safety ranking or a general collision-avoidance claim. TTC would additionally depend on validated speed, timing and closing-sign conventions.')
        p('No fuel flow, electrical power or calibrated energy model is recorded. Fuel consumption, Wh/km and percentage energy savings therefore cannot be established. The earlier positive-kinetic-variation diagnostic measures repeated speed increases and excludes road load, efficiency and regeneration; it is not a fuel result.')

        page(); h('4 Discussion and conclusions')
        h('What the look ahead information contributes', 2)
        p('An upstream leader can reveal a disturbance before it fully appears in the middle vehicle\'s speed. The known middle IDM gives PID and LQI a basis for anticipating that response; MPC uses predecessor trends. Feedforward can therefore complement immediate spacing feedback, as motivated by cooperative-control studies [1-3]. Hidden command memory, asynchronous telemetry, safety overrides and response lag still limit prediction.')
        p('Earlier slowing can reduce a later speed transient while enlarging the gap temporarily. Feedback must subsequently recover that separation. This explains why spacing accuracy and attenuation can trade off, but the physical tests do not isolate this mechanism: every candidate uses look-ahead information, MPC has a shorter policy, and there is no matched physical controller using only the immediate predecessor.')
        h('Conclusions supported by the records', 2)
        p('The MPC-labelled run maintains the smallest gaps under its chosen policy. PID shows the lowest student/middle reported-speed variation ratios in the two excited nonlaunch events. LQI recovers substantially from its larger event-3 entry gap. All three retain positive recorded active gaps, with no active spacing-shield flags. The records support these scenario-specific observations and do not select one method across tracking, safety, comfort, stability and energy.')
        p('Any future causal comparison would need common targets, initial conditions and disturbances, repeated runs and explicit optimizer-status logging. A matched look-ahead ablation would separate the contribution of upstream information from controller tuning. Those are limits of the current evidence; the completed experiments and saved results are retained here without replacement.')
        h('References', 2)
        p('[1] G. J. L. Naus et al. Cooperative adaptive cruise control, design and experiments. American Control Conference, 2010, pp. 6145-6150. DOI: 10.1109/ACC.2010.5531596.')
        p('[2] C. Wang et al. Cooperative adaptive cruise control for connected autonomous vehicles by factoring communication-related constraints. Transportation Research Part C 113, 2020, pp. 124-145. DOI: 10.1016/j.trc.2019.04.010.')
        p('[3] C. Wu, Y. Lin and A. Eskandarian. Cooperative Adaptive Cruise Control With Adaptive Kalman Filter Subject to Temporary Communication Loss. IEEE Access, 2019. DOI: 10.1109/ACCESS.2019.2928004.')
        h('Experimental record', 2)
        p('Sources are the Lab II guidelines, supplied experiment/controller files, the three Group 4 raw-run folders, and saved physical_summary.json, event CSVs, macro_metrics.json and closed_loop.json in results/three_controller_analysis. The earlier deployment-error record is retained separately. Controller code and original data/results are unchanged; report tables use recorded observations or explicitly identified historical simulations.')


        def escape(s):
            for a,b in [('\\', r'\textbackslash{}'), ('&',r'\&'), ('%',r'\%'), ('_',r'\_'), ('#',r'\#')]:
                s=s.replace(a,b)
            return s.replace('γ',r'$\gamma$').replace('β',r'$\beta$').replace('−','-')


        tex = [r'''\documentclass[11pt,a4paper]{article}
        \usepackage[margin=20mm]{geometry}
        \usepackage[T1]{fontenc}
        \usepackage{lmodern}
        \usepackage{amsmath,amssymb,graphicx,booktabs,float,array,caption}
        \usepackage[hidelinks]{hyperref}
        \setlength{\parindent}{0pt}
        \setlength{\parskip}{5pt}
        \setlength{\tabcolsep}{5pt}
        \emergencystretch=2em
        \renewcommand{\arraystretch}{1.12}
        \captionsetup{font=small,labelfont=bf}
        \begin{document}''']

        # Start from the retained draft to preserve package/theme relationships; keep
        # the original file untouched. Explicit style alignment supersedes its geometry.
        ref=ROOT/'report/archive/lab_02_report_draft_20261020.docx'
        original_bytes=ref.read_bytes()
        doc=Document(io.BytesIO(original_bytes))
        sect=deepcopy(doc.sections[0]._sectPr)
        for child in list(doc._element.body): doc._element.body.remove(child)
        doc._element.body.append(sect)
        for sec in doc.sections:
            sec.page_width=Mm(210);sec.page_height=Mm(297)
            sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(20)
        for name,size in [('Normal',11),('Title',20),('Subtitle',12),('Heading 1',14),('Heading 2',11.5),('Caption',9.5)]:
            st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size)
            st.font.color.rgb=RGBColor(0,0,0)
            st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.04
            st.paragraph_format.keep_with_next=name.startswith('Heading')
            if name.startswith('Heading'):
                st.font.bold=True;st.paragraph_format.space_before=Pt(8)
        doc.styles['Title'].font.bold=True

        pending_break=False
        for block in blocks:
            kind=block[0]
            if kind=='title':
                tex.append(r'\begin{center}{\LARGE\bfseries '+escape(block[1])+r'}\\[5pt]')
                par=doc.add_paragraph(block[1], 'Title');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
            elif kind=='subtitle':
                tex.append(r'{\large INTR 5100 \quad Lab Session II}\\[5pt]Group 4 \quad HKUST(GZ) \quad Fall 2026\end{center}')
                par=doc.add_paragraph(block[1]+'\nHKUST(GZ)   Fall 2026','Subtitle');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
            elif kind=='page':
                tex.append(r'\newpage');pending_break=True
            elif kind=='h':
                level,text=block[1:]
                tex.append(('\\section*{' if level==1 else '\\subsection*{')+escape(text)+'}')
                par=doc.add_paragraph(text,f'Heading {level}')
                if pending_break:
                    par.paragraph_format.page_break_before=True
                    pending_break=False
            elif kind=='p':
                tex.append(escape(block[1])+'\n');doc.add_paragraph(block[1])
            elif kind=='equation':
                tex.append('\\[\n'+block[1]+'\n\\]')
                par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
                math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=block[2]
                r.append(t);math.append(r);par._p.append(math)
            elif kind=='table':
                caption,headers,values,widths=block[1:]
                alignment='l'+'r'*(len(headers)-1) if len(headers)>2 else 'll'
                tex.append(r'\begin{table}[H]\centering\small\caption{'+escape(caption)+'}'+
                           '\n\\begin{tabular}{'+alignment+'}\n\\toprule\n'+
                           ' & '.join(escape(s) for s in headers)+r' \\'+ '\n\\midrule')
                tex.extend(' & '.join(escape(str(s)) for s in row)+r' \\' for row in values)
                tex.append(r'\bottomrule\end{tabular}\end{table}')
                cap=doc.add_paragraph(caption,'Caption');cap.paragraph_format.keep_with_next=True
                tab=doc.add_table(rows=1,cols=len(headers));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
                # Scale reference widths to the aligned 170-mm text area.
                widths=[w/sum(widths)*170/25.4 for w in widths]
                for c,w in zip(tab.columns,widths): c.width=Inches(w)
                for row_values in [headers]+values:
                    row=tab.rows[0] if row_values is headers else tab.add_row()
                    for j,(cell,value,w) in enumerate(zip(row.cells,row_values,widths)):
                        cell.width=Inches(w);cell.text=str(value);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                        for side in ('top','left','bottom','right'):
                            edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                        pr.append(borders)
                        mar=OxmlElement('w:tcMar')
                        for side in ('top','left','bottom','right'):
                            e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');mar.append(e)
                        pr.append(mar)
                        for pp in cell.paragraphs:
                            pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.line_spacing=1.0
                            pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(headers)==2 else WD_ALIGN_PARAGRAPH.CENTER
                            for run in pp.runs:run.font.size=Pt(9.5);run.bold=row_values is headers
                        if row_values is headers:
                            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');pr.append(sh)
                    no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
                doc.add_paragraph().paragraph_format.space_after=Pt(0)
            elif kind=='figure':
                stem,caption=block[1:]
                tex.append(r'\begin{figure}[H]\centering\includegraphics[width=\linewidth]{figures/'+stem+r'.pdf}\caption{'+escape(caption)+r'}\end{figure}')
                par=doc.add_paragraph();par.paragraph_format.keep_with_next=True
                par.alignment=WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(str(FIG/(stem+'.png')),width=Mm(150))
                doc.add_paragraph(caption,'Caption')

        tex.append(r'\end{document}')
        (ROOT/'report/source/main.tex').write_text('\n\n'.join(tex))
        docx=ROOT/'report/lab_02_report.docx'
        doc.save(docx)
        # Preserve opaque parts of the reference package, including the retained theme.
        with ZipFile(io.BytesIO(original_bytes)) as z: original={n:z.read(n) for n in z.namelist()}
        with ZipFile(docx) as z: built={n:z.read(n) for n in z.namelist()}
        editable={'word/document.xml','word/styles.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
        for n,b in original.items():
            if n not in editable: built[n]=b
        with ZipFile(docx,'w',ZIP_DEFLATED) as z:
            for n,b in built.items():z.writestr(n,b)
        with ZipFile(docx) as z:
            for n,b in original.items():
                if n not in editable:assert z.read(n)==b
        assert ref.read_bytes()==original_bytes
        after=protected();assert before==after
        (ROOT/'review/report_build_checks.json').write_text(json.dumps({
            'scope':'Report-only revision; no fitting, controller testing or simulation was run.',
            'protected_file_count':len(before),'protected_files_unchanged':before==after,
            'original_draft_unchanged':True,'checks':checks,'protected_sha256':before,
            'saved_simulation_count':len(closed['runs'])},indent=2))
        print(f'Created aligned DOCX and LaTeX report; {len(before)} original files unchanged.')

        if not args.skip_pdf:
            compiler = shutil.which('tectonic') or str(Path(sys.executable).with_name('tectonic'))
            with tempfile.TemporaryDirectory(prefix='intr5100-lab2-tex-') as build_dir:
                command = [compiler, '--outdir', build_dir, 'main.tex']
                if args.offline:
                    command.insert(1, '--only-cached')
                subprocess.run(command, cwd=ROOT/'report/source', check=True)
                shutil.copyfile(Path(build_dir)/'main.pdf', ROOT/'report/lab_02_report.pdf')
            print('Compiled report/lab_02_report.pdf; temporary compiler files removed.')


if __name__ == '__main__':
    main()
